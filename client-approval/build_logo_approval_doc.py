from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\02-239241-036\Documents\GitHub\DHA-Karachi-Villas-Website")
IMAGE_ROOT = Path(r"C:\Users\02-239241-036\.codex\generated_images\01a124f4-bbb6-7402-85b0-38446ee15d28")
OUTPUT = ROOT / "client-approval" / "DHA-Karachi-Villas-Logo-Approval-Options.docx"

NAVY = RGBColor(20, 43, 58)
CORAL = RGBColor(200, 59, 77)
MUTED = RGBColor(82, 98, 110)
BLACK = RGBColor(0, 0, 0)

OPTIONS = {
    "current": [
        ("Current Identity Refined", "exec-3a92f6bb-8eb0-4f4b-9649-ca8bb6dca034.png", "Thicker house outlines for stronger visibility in the website header and social media use."),
    ],
    "monogram": [
        ("DV Monogram Original", "exec-744eef3f-19cb-41dd-995e-5b48f8f72a54.png", "Clean DV lettermark with a modern villa roof accent."),
        ("DV Monogram Compact", "exec-888b6092-7dc9-4f7b-a06b-ba1f27063ede.png", "Compact roof-and-doorway variation designed to work well as a favicon and app icon."),
        ("DV Monogram Gateway", "exec-c7aafe11-9a4b-4a0e-a500-09e297245e35.png", "Bold interlocking DV mark with a premium architectural gateway feel."),
    ],
    "villa": [
        ("DV Monogram Ribbon", "exec-12e130d8-8bc3-4935-855b-0a871ce32df9.png", "Softer, more editorial DV mark with a subtle doorway accent."),
        ("Villa Location Emblem", "exec-c1cd9f22-9195-4998-8d4b-84207a7c8dca.png", "A villa and location-pin concept for a more descriptive real-estate mark."),
        ("Villa Gate Emblem", "exec-dca93f46-3848-479d-a005-2325395a39eb.png", "A strong residential gate and villa icon with a luxury property feel."),
    ],
}


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def style_run(run, size=None, bold=None, color=None, font="Aptos"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_text(doc, text, size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    style_run(r, size=size, color=color)
    r.italic = italic
    return p


def add_heading(doc, text, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    style_run(r, size=18, bold=True, color=BLACK)
    return p


def add_option(doc, number, title, filename, note):
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(8)
    label.paragraph_format.space_after = Pt(3)
    run = label.add_run(f"{number:02d}  {title}")
    style_run(run, size=11.5, bold=True, color=BLACK)

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_paragraph.add_run().add_picture(str(IMAGE_ROOT / filename), width=Inches(2.05))

    note_paragraph = doc.add_paragraph()
    note_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_paragraph.paragraph_format.space_after = Pt(10)
    note_run = note_paragraph.add_run(note)
    style_run(note_run, size=9.3, color=MUTED)


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    add_page_number(paragraph)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(40)
    title.paragraph_format.space_after = Pt(12)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DHA Karachi Villas Logo Approval Options")
    style_run(run, size=25, bold=True, color=BLACK, font="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Brand identity options prepared for client selection")
    style_run(run, size=12, color=MUTED)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(18)
    run = intro.add_run("This document brings together the shortlisted logo directions for DHA Karachi Villas. Please select one preferred main logo for final production.")
    style_run(run, size=10.5, color=MUTED)

    current_title, current_file, current_note = OPTIONS["current"][0]
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(5)
    run = label.add_run(current_title)
    style_run(run, size=13, bold=True, color=BLACK)

    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_after = Pt(8)
    logo.add_run().add_picture(str(IMAGE_ROOT / current_file), width=Inches(3.05))

    add_text(doc, current_note, size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    add_text(doc, "The current refined logo is shown above. The following pages contain six alternative directions for comparison before final selection.", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, italic=True)

    pages = [
        ("DV Monogram Concepts", "Compact DV concepts designed for a premium modern brand and strong small-size performance.", OPTIONS["monogram"]),
        ("Additional Architectural Concepts", "Alternative logo directions that combine villa architecture with a distinct brand mark.", OPTIONS["villa"]),
    ]

    option_number = 1
    for title_text, description, options in pages:
        add_heading(doc, title_text, page_break=True)
        add_text(doc, description, size=10, color=MUTED, space_after=6)
        for option in options:
            add_option(doc, option_number, *option)
            option_number += 1

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
